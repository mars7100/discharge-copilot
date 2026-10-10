"""Structured draft exports with source appendix (spec §10 Export output)."""
from datetime import datetime
from io import BytesIO

from docx import Document as WordDocument
from docx.shared import Inches, Pt

from core.models import TEMPLATE, ClinicianClarification, Document, Draft, Resolution, SourceSpan

SECTIONS = {s["section_id"]: s["title"] for s in TEMPLATE}
NOT_DOCUMENTED = "Not documented in the supplied notes."


def marker(span: SourceSpan) -> str:
    return f"[N{span.document_id[-2:]}:B{span.order:02d}]"


def to_markdown(draft: Draft, spans: dict[str, SourceSpan], documents: list[Document], case_id: str = "case_001",
                clarifications: list[ClinicianClarification] = (), resolutions: list[Resolution] = (),
                missing_notes: list[str] = ()) -> str:
    status = (f"APPROVED by {draft.reviewer} at {draft.approved_at:%Y-%m-%d %H:%M}" if draft.status == "approved"
              else "DRAFT — NOT APPROVED")
    out = ["# Discharge summary", "", "**SYNTHETIC DEMONSTRATION — NOT FOR CLINICAL USE**", "",
           f"- Case: {case_id} (fictional)", f"- Status: {status}", f"- Generated: {datetime.now():%Y-%m-%d %H:%M}", ""]
    if missing_notes:
        out += [f"**Incomplete: notes not included (extraction failed): {', '.join(missing_notes)}**", ""]
    for section in TEMPLATE:
        claims = [c for c in draft.claims if c.section == section["section_id"]]
        if not claims and section["optional"]:
            continue
        out += [f"## {section['title']}", ""]
        claims.sort(key=lambda c: c.problem is not None)  # corrections without a problem go first, no subheading
        for i, c in enumerate(claims):
            if c.problem and (i == 0 or c.problem != claims[i - 1].problem):
                out.append(f"### {c.problem}")
            flag = " **(needs review)**" if c.support_status == "needs_review" else ""
            clar = " *(clinician-provided)*" if c.clinician_clarification_ids else ""
            out.append(f"- {c.text}{flag}{clar} " + " ".join(marker(spans[s]) for s in c.evidence_span_ids))
        out += ([] if claims else [NOT_DOCUMENTED]) + [""]
    out += ["## Unresolved discrepancies", ""]
    resolved = {r.conflict_id for r in resolutions}
    open_ = [d for d in draft.discrepancies if d.discrepancy_id not in resolved]  # corrections alone do not resolve one
    for d in open_:
        out.append(f"- {d.discrepancy_id} ({d.status}): {d.description} "
                   + " ".join(marker(spans[s]) for s in d.source_span_ids))
    out += ([] if open_ else ["None flagged."]) + ["", "## Conflict resolutions", ""]
    desc = {d.discrepancy_id: d.description for d in draft.discrepancies}
    out += [f"- {r.conflict_id}: {desc.get(r.conflict_id, '')} — \"{r.text}\" (clinician-provided; "
            + (f"chosen from {marker(spans[r.chosen_span_id])}" if r.chosen_span_id else "typed clarification") + ") "
            + " ".join(marker(spans[q.span_id]) for q in r.quotes) for r in resolutions] or ["None."]
    out += ["", "## Clinician-provided clarifications", ""]
    accepted = [k for k in clarifications if k.status == "accepted"]
    out += [f"- {k.clarification_id} → {k.target_claim_id}: \"{k.statement}\" (clinician-provided, not document evidence)"
            for k in accepted] or ["None."]
    out += ["", "## Source appendix", ""]
    files = {d.document_id: d for d in documents}
    cited = {s for c in draft.claims for s in c.evidence_span_ids} | {s for d in draft.discrepancies for s in d.source_span_ids} \
        | {q.span_id for r in resolutions for q in r.quotes}
    for sid in sorted(cited):
        sp = spans[sid]
        doc = files.get(sp.document_id)
        when = f", {doc.date} (Hospital Day {doc.hospital_day})" if doc else ""
        out += [f"**{marker(sp)}** {sp.filename}, page {sp.page}{when}", "", "> " + sp.text.replace("\n", "\n> "), ""]
    return "\n".join(out)


def to_docx(draft: Draft, spans: dict[str, SourceSpan], documents: list[Document], case_id: str = "case_001",
            clarifications: list[ClinicianClarification] = (), resolutions: list[Resolution] = (),
            missing_notes: list[str] = ()) -> bytes:
    """Render draft content and provenance as editable Word paragraphs and runs."""
    word = WordDocument()
    normal = word.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for section in word.sections:
        section.top_margin = section.bottom_margin = Inches(1)
        section.left_margin = section.right_margin = Inches(1)
    word.add_heading("Discharge summary", level=0)
    word.add_paragraph().add_run("SYNTHETIC DEMONSTRATION — NOT FOR CLINICAL USE").bold = True
    status = (f"APPROVED by {draft.reviewer} at {draft.approved_at:%Y-%m-%d %H:%M}" if draft.status == "approved"
              else "DRAFT — NOT APPROVED")
    for text in (f"Case: {case_id} (fictional)", f"Status: {status}",
                 f"Generated: {datetime.now():%Y-%m-%d %H:%M}"):
        word.add_paragraph(text)
    if missing_notes:
        word.add_paragraph().add_run(
            f"Incomplete: notes not included (extraction failed): {', '.join(missing_notes)}").bold = True
    for section in TEMPLATE:
        claims = sorted((c for c in draft.claims if c.section == section["section_id"]),
                        key=lambda c: c.problem is not None)
        if not claims and section["optional"]:
            continue
        word.add_heading(section["title"], level=1)
        if not claims:
            word.add_paragraph(NOT_DOCUMENTED)
        for i, claim in enumerate(claims):
            if claim.problem and (i == 0 or claim.problem != claims[i - 1].problem):
                word.add_heading(claim.problem, level=2)
            paragraph = word.add_paragraph(claim.text, style="List Bullet")
            if claim.support_status == "needs_review":
                paragraph.add_run(" (needs review)").bold = True
            if claim.clinician_clarification_ids:
                paragraph.add_run(" (clinician-provided)").italic = True
            if claim.evidence_span_ids:
                paragraph.add_run(" " + " ".join(marker(spans[s]) for s in claim.evidence_span_ids))
    word.add_heading("Unresolved discrepancies", level=1)
    resolved = {r.conflict_id for r in resolutions}
    open_discrepancies = [d for d in draft.discrepancies if d.discrepancy_id not in resolved]
    for discrepancy in open_discrepancies:
        word.add_paragraph(f"{discrepancy.discrepancy_id} ({discrepancy.status}): {discrepancy.description} "
                           + " ".join(marker(spans[s]) for s in discrepancy.source_span_ids), style="List Bullet")
    if not open_discrepancies:
        word.add_paragraph("None flagged.")
    word.add_heading("Conflict resolutions", level=1)
    descriptions = {d.discrepancy_id: d.description for d in draft.discrepancies}
    for resolution in resolutions:
        origin = (f"chosen from {marker(spans[resolution.chosen_span_id])}" if resolution.chosen_span_id
                  else "typed clarification")
        word.add_paragraph(
            f'{resolution.conflict_id}: {descriptions.get(resolution.conflict_id, "")} — "{resolution.text}" '
            f"(clinician-provided; {origin}) " + " ".join(marker(spans[q.span_id]) for q in resolution.quotes),
            style="List Bullet")
    if not resolutions:
        word.add_paragraph("None.")
    word.add_heading("Clinician-provided clarifications", level=1)
    accepted = [c for c in clarifications if c.status == "accepted"]
    for clarification in accepted:
        word.add_paragraph(
            f'{clarification.clarification_id} → {clarification.target_claim_id}: "{clarification.statement}" '
            "(clinician-provided, not document evidence)", style="List Bullet")
    if not accepted:
        word.add_paragraph("None.")
    word.add_heading("Source appendix", level=1)
    files = {d.document_id: d for d in documents}
    cited = ({s for c in draft.claims for s in c.evidence_span_ids}
             | {s for d in draft.discrepancies for s in d.source_span_ids}
             | {q.span_id for r in resolutions for q in r.quotes}
             | {r.chosen_span_id for r in resolutions if r.chosen_span_id})
    for sid in sorted(cited):
        span = spans[sid]
        document = files.get(span.document_id)
        when = f", {document.date} (Hospital Day {document.hospital_day})" if document else ""
        paragraph = word.add_paragraph()
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run(marker(span)).bold = True
        paragraph.add_run(f" {span.filename}, page {span.page}{when}")
        for line in span.text.split("\n"):
            word.add_paragraph(line, style="Quote")
    output = BytesIO()
    word.save(output)
    return output.getvalue()
