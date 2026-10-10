from datetime import datetime
from io import BytesIO

from docx import Document as WordDocument

from core.export import NOT_DOCUMENTED, to_docx, to_markdown
from core.models import ClinicianClarification, Discrepancy, Document, Draft, EvidenceRef, Resolution, SourceSpan, SummaryClaim

SPAN = SourceSpan(span_id="note_04_p01_b07", document_id="note_04", filename="note_04.txt",
                  page=1, order=7, text="Medication A 10 mg daily.")


def test_empty_required_sections_say_not_documented_and_empty_optional_ones_are_omitted():
    claim = SummaryClaim(claim_id="C01", section="discharge_medications", text="Medication A 10 mg daily.",
                         evidence_span_ids=[SPAN.span_id], support_status="unverified")
    md = to_markdown(Draft(claims=[claim]), {SPAN.span_id: SPAN}, [])
    headings = [line for line in md.splitlines() if line.startswith("## ")][:5]
    assert headings == ["## Reason for admission", "## Hospital course", "## Condition at discharge",
                        "## Discharge medications", "## Follow-up"]
    assert md.count(NOT_DOCUMENTED) == 4


def test_docx_preserves_structured_content_provenance_and_editable_styles() -> None:
    """Reopen actual Word bytes and check ordering, flags, resolutions, and source text."""
    claim = SummaryClaim(claim_id="C01", section="hospital_course", text="Problem claim <literal> & text.",
                         problem="Pneumonia", evidence_span_ids=[SPAN.span_id], support_status="needs_review",
                         clinician_clarification_ids=["K01"])
    correction = claim.model_copy(update={"claim_id": "C02", "problem": None, "text": "Correction first."})
    draft = Draft(claims=[claim, correction], status="approved", reviewer="Dr Test",
                  approved_at=datetime(2026, 1, 2, 12, 30), discrepancies=[
                      Discrepancy(discrepancy_id=id_, description=f"Conflict {id_}", source_span_ids=[SPAN.span_id])
                      for id_ in ("D01", "D02", "D03")])
    source = SPAN.model_copy(update={"text": "Exact first line.\nExact second line."})
    document = Document(document_id=SPAN.document_id, filename=SPAN.filename, date="2026-01-01",
                        hospital_day=2, type="Note", author_role=None)
    resolutions = [Resolution(conflict_id=id_, text=text, chosen_span_id=chosen,
                              quotes=[EvidenceRef(span_id=SPAN.span_id, quote="Exact first line.")])
                   for id_, text, chosen in (("D01", "Selected value", SPAN.span_id), ("D02", "Typed plan", None))]
    clarifications = [ClinicianClarification(clarification_id=id_, target_claim_id="C01", statement=text, status=status)
                      for id_, text, status in (("K01", "Accepted wording", "accepted"),
                                               ("K02", "Rejected wording", "rejected"),
                                               ("K03", "Pending wording", "pending"))]
    word = WordDocument(BytesIO(to_docx(draft, {SPAN.span_id: source}, [document], "case_test",
                                       clarifications, resolutions, ["failed_note.txt"])))
    paragraphs = word.paragraphs
    text = "\n".join(p.text for p in paragraphs)
    assert "SYNTHETIC DEMONSTRATION — NOT FOR CLINICAL USE" in text
    assert "Case: case_test (fictional)" in text
    assert "APPROVED by Dr Test at 2026-01-02 12:30" in text and "Generated:" in text
    assert "Incomplete: notes not included (extraction failed): failed_note.txt" in text
    assert text.index("Correction first.") < text.index("Pneumonia") < text.index(claim.text)
    assert f"{claim.text} (needs review) (clinician-provided) [N04:B07]" in text
    unresolved = text.split("Unresolved discrepancies")[1].split("Conflict resolutions")[0]
    assert "D03 (open): Conflict D03 [N04:B07]" in unresolved
    assert "D01" not in unresolved and "D02" not in unresolved
    assert '"Selected value" (clinician-provided; chosen from [N04:B07]) [N04:B07]' in text
    assert '"Typed plan" (clinician-provided; typed clarification) [N04:B07]' in text
    assert 'K01 → C01: "Accepted wording" (clinician-provided, not document evidence)' in text
    assert "Rejected wording" not in text and "Pending wording" not in text
    assert "[N04:B07] note_04.txt, page 1, 2026-01-01 (Hospital Day 2)" in text
    assert "Exact first line.\nExact second line." in text
    headings = [p.text for p in paragraphs if p.style.name == "Heading 1"]
    assert headings[:4] == ["Reason for admission", "Hospital course", "Condition at discharge", "Follow-up"]
    assert "Discharge medications" not in headings
    assert text.count(NOT_DOCUMENTED) == 3
    assert next(p for p in paragraphs if p.text == "Pneumonia").style.name == "Heading 2"
    claim_paragraph = next(p for p in paragraphs if p.text.startswith(claim.text))
    assert claim_paragraph.style.name == "List Bullet"
    assert claim_paragraph.runs[1].bold and claim_paragraph.runs[2].italic
    assert word.styles["Normal"].font.name == "Calibri"
    claim_paragraph.add_run(" Editable addition.")
    output = BytesIO()
    word.save(output)
    assert "Editable addition." in "\n".join(p.text for p in WordDocument(BytesIO(output.getvalue())).paragraphs)


def test_docx_empty_draft_has_required_sections_and_empty_provenance() -> None:
    """Empty optional sections disappear while required sections remain explicit."""
    word = WordDocument(BytesIO(to_docx(Draft(claims=[]), {}, [])))
    text = "\n".join(p.text for p in word.paragraphs)
    assert text.count(NOT_DOCUMENTED) == 4
    assert "DRAFT — NOT APPROVED" in text
    assert "None flagged." in text and text.count("None.") == 2
